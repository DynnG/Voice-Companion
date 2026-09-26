"""
Document Text Extraction Service for Pal Voice Companion.
Extracts clean plain text from PDF, DOCX, TXT, and Markdown files.
"""

import io
import os
import logging
from typing import Dict, Any, Optional

logger = logging.getLogger("voice-companion-backend.document-service")

try:
    import pypdf
except ImportError:
    pypdf = None
    logger.warning("pypdf is not installed. PDF text extraction will be unavailable.")

try:
    import docx
except ImportError:
    docx = None
    logger.warning("python-docx is not installed. DOCX text extraction will be unavailable.")


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract plain text from PDF file bytes using pypdf."""
    if not pypdf:
        raise RuntimeError("pypdf library is required for PDF text extraction.")
    
    stream = io.BytesIO(file_bytes)
    reader = pypdf.PdfReader(stream)
    pages_text = []
    
    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                pages_text.append(page_text.strip())
        except Exception as e:
            logger.warning(f"Error extracting text from PDF page {idx + 1}: {e}")
            continue

    return "\n\n".join(pages_text).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract plain text from DOCX file bytes using python-docx."""
    if not docx:
        raise RuntimeError("python-docx library is required for DOCX text extraction.")
    
    stream = io.BytesIO(file_bytes)
    doc = docx.Document(stream)
    lines = []

    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines).strip()


def extract_plain_text(file_bytes: bytes) -> str:
    """Extract plain text from TXT or MD file bytes trying common encodings."""
    encodings = ("utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1")
    for enc in encodings:
        try:
            return file_bytes.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()


def extract_document_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Main extraction dispatcher for PDF, DOCX, TXT, and MD files.
    Returns dictionary with extracted_text, char_count, file_type, and status.
    """
    if not file_bytes:
        return {
            "filename": filename,
            "file_type": "unknown",
            "extracted_text": "",
            "char_count": 0,
            "status": "empty_file"
        }

    ext = os.path.splitext(filename)[1].lower().lstrip(".")

    try:
        if ext == "pdf":
            text = extract_pdf_text(file_bytes)
        elif ext in ("docx", "doc"):
            text = extract_docx_text(file_bytes)
        elif ext in ("txt", "md", "markdown", "text"):
            text = extract_plain_text(file_bytes)
        else:
            # Fallback attempt as plain text
            text = extract_plain_text(file_bytes)

        return {
            "filename": filename,
            "file_type": ext,
            "extracted_text": text,
            "char_count": len(text),
            "status": "success"
        }
    except Exception as e:
        logger.error(f"Failed to extract text from {filename} ({ext}): {e}", exc_info=True)
        return {
            "filename": filename,
            "file_type": ext,
            "extracted_text": "",
            "char_count": 0,
            "status": "error",
            "error": str(e)
        }
